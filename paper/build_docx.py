"""Builds the Indonesian paper draft: paper/draft_paper_id.docx

Usage: uv run --with python-docx python paper/build_docx.py
Numbers in the accuracy table are read from paper/tables/*.csv (regenerate with `python -m src.analysis`);
the remaining numbers are copied from RESULTS.md.
"""
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "draft_paper_id.docx"  # v0.1 archived as draft_paper_id_v0.1_pre-humanizer.docx
FIG, TAB = ROOT / "figures", ROOT / "tables"

# ---------------------------------------------------------------- references (IEEE style, numbered by first citation)
REFS = {
    "shahrad": 'M. Shahrad, R. Fonseca, Í. Goiri, G. Chaudhry, P. Batum, J. Cooke, E. Laureano, C. Tresness, M. Russinovich, dan R. Bianchini, "Serverless in the wild: Characterizing and optimizing the serverless workload at a large cloud provider," dalam Proc. USENIX ATC, 2020, hlm. 205–218.',
    "roy": 'N. Roy, A. Dubey, dan A. Gokhale, "Efficient autoscaling in the cloud using predictive models for workload forecasting," dalam Proc. IEEE CLOUD, 2011, hlm. 500–507, doi: 10.1109/CLOUD.2011.42.',
    "hpa": 'The Kubernetes Authors, "Horizontal Pod Autoscaling," dokumentasi Kubernetes. [Daring]. Tersedia: https://kubernetes.io/docs/tasks/run-application/horizontal-pod-autoscale/ (diakses 25 Oktober 2026).',
    "chronos": 'A. F. Ansari dkk., "Chronos: Learning the language of time series," Transactions on Machine Learning Research, 2024.',
    "timesfm": 'A. Das, W. Kong, R. Sen, dan Y. Zhou, "A decoder-only foundation model for time-series forecasting," dalam Proc. ICML, PMLR vol. 235, 2024, hlm. 10148–10167.',
    "moirai": 'G. Woo, C. Liu, A. Kumar, C. Xiong, S. Savarese, dan D. Sahoo, "Unified training of universal time series forecasting transformers," dalam Proc. ICML, PMLR vol. 235, 2024, hlm. 53140–53164.',
    "ttm": 'V. Ekambaram, A. Jati, P. Dayama, S. Mukherjee, N. H. Nguyen, W. M. Gifford, C. Reddy, dan J. Kalagnanam, "Tiny Time Mixers (TTMs): Fast pre-trained models for enhanced zero/few-shot forecasting of multivariate time series," dalam Proc. NeurIPS, 2024.',
    "toner": 'W. Toner, T. L. Lee, A. Joosen, R. Singh, dan M. Asenov, "Performance of zero-shot time series foundation models on cloud data," ICLR 2025 Workshop "I Can\'t Believe It\'s Not Better", arXiv:2502.12944, 2025.',
    "chronos2": 'A. F. Ansari dkk., "Chronos-2: From univariate to universal forecasting," arXiv:2510.15821, 2025.',
    "timesfm3": 'A. Jain dan R. Sen, "TimesFM-3: A zero-shot foundation model for multivariate forecasting," Google Research Blog, 31 Agustus 2026. [Daring]. Tersedia: https://research.google/blog/timesfm-3-a-zero-shot-foundation-model-for-multivariate-forecasting/',
    "timesfm25card": 'Google Research, "timesfm-2.5-200m-pytorch," model card Hugging Face, 2025. [Daring]. Tersedia: https://huggingface.co/google/timesfm-2.5-200m-pytorch',
    "moirai2": 'C. Liu, T. Aksu, J. Liu, X. Liu, H. Yan, Q. Pham, S. Savarese, D. Sahoo, C. Xiong, dan J. Li, "Moirai 2.0: When less is more for time series forecasting," arXiv:2511.11698, 2025.',
    "tirex": 'A. Auer, P. Podest, D. Klotz, S. Böck, G. Klambauer, dan S. Hochreiter, "TiRex: Zero-shot forecasting across long and short horizons with enhanced in-context learning," dalam Proc. NeurIPS, 2025.',
    "toto": 'B. Cohen dkk., "This time is different: An observability perspective on time series foundation models," arXiv:2505.14766, 2025.',
    "gifteval": 'T. Aksu, G. Woo, J. Liu, X. Liu, C. Liu, S. Savarese, C. Xiong, dan D. Sahoo, "GIFT-Eval: A benchmark for general time series forecasting model evaluation," arXiv:2410.10393, 2024.',
    "cloudops": 'G. Woo, C. Liu, A. Kumar, dan D. Sahoo, "Pushing the limits of pre-training for time series forecasting in the CloudOps domain," arXiv:2310.05063, 2023.',
    "zhang21": 'Y. Zhang, Í. Goiri, G. I. Chaudhry, R. Fonseca, S. Elnikety, C. Delimitrou, dan R. Bianchini, "Faster and cheaper serverless computing on harvested resources," dalam Proc. ACM SOSP, 2021, hlm. 724–739, doi: 10.1145/3477132.3483580.',
    "joosen": 'A. Joosen, A. Hassan, M. Asenov, R. Singh, L. Darlow, J. Wang, dan A. Barker, "How does it function? Characterizing long-term trends in production serverless workloads," dalam Proc. ACM SoCC, 2023, hlm. 443–458.',
    "guo": 'J. Guo, Z. Chang, S. Wang, H. Ding, Y. Feng, L. Mao, dan Y. Bao, "Who limits the resource efficiency of my datacenter: An analysis of Alibaba datacenter traces," dalam Proc. IEEE/ACM IWQoS, 2019, hlm. 1–10, doi: 10.1145/3326285.3329074.',
    "syntetos": 'A. A. Syntetos, J. E. Boylan, dan J. D. Croston, "On the categorization of demand patterns," Journal of the Operational Research Society, vol. 56, no. 5, hlm. 495–503, 2005.',
    "mstl": 'K. Bandara, R. J. Hyndman, dan C. Bergmeir, "MSTL: A seasonal-trend decomposition algorithm for time series with multiple seasonal patterns," International Journal of Operational Research, vol. 52, no. 1, hlm. 79–98, 2025, doi: 10.1504/IJOR.2025.143957.',
    "croston": 'J. D. Croston, "Forecasting and stock control for intermittent demands," Operational Research Quarterly, vol. 23, no. 3, hlm. 289–303, 1972.',
    "tsb": 'R. H. Teunter, A. A. Syntetos, dan M. Z. Babai, "Intermittent demand: Linking forecasting to inventory obsolescence," European Journal of Operational Research, vol. 214, no. 3, hlm. 606–615, 2011.',
    "statsforecast": 'A. Garza, M. Mergenthaler Canseco, C. Challú, dan K. G. Olivares, "StatsForecast: Lightning fast forecasting with statistical and econometric models," PyCon Salt Lake City, 2022.',
    "lightgbm": 'G. Ke, Q. Meng, T. Finley, T. Wang, W. Chen, W. Ma, Q. Ye, dan T.-Y. Liu, "LightGBM: A highly efficient gradient boosting decision tree," dalam Proc. NeurIPS, 2017.',
    "nbeats": 'B. N. Oreshkin, D. Carpov, N. Chapados, dan Y. Bengio, "N-BEATS: Neural basis expansion analysis for interpretable time series forecasting," dalam Proc. ICLR, 2020.',
    "lstm": 'S. Hochreiter dan J. Schmidhuber, "Long short-term memory," Neural Computation, vol. 9, no. 8, hlm. 1735–1780, 1997.',
    "neuralforecast": 'K. G. Olivares, C. Challú, A. Garza, M. Mergenthaler Canseco, dan A. Dubrawski, "NeuralForecast: User friendly state-of-the-art neural forecasting models," PyCon Salt Lake City, 2022.',
    "hyndman": 'R. J. Hyndman dan A. B. Koehler, "Another look at measures of forecast accuracy," International Journal of Forecasting, vol. 22, no. 4, hlm. 679–688, 2006.',
    "vovk": 'V. Vovk, A. Gammerman, dan G. Shafer, Algorithmic Learning in a Random World. New York: Springer, 2005.',
    "demsar": 'J. Demšar, "Statistical comparisons of classifiers over multiple data sets," Journal of Machine Learning Research, vol. 7, hlm. 1–30, 2006.',
    "golec": 'M. Golec, G. K. Walia, M. Kumar, F. Cuadrado, S. S. Gill, dan S. Uhlig, "Cold start latency in serverless computing: A systematic review, taxonomy, and future directions," ACM Computing Surveys, vol. 57, no. 3, terbit daring 11 November 2024, doi: 10.1145/3700875.',
    "jegannathan": 'A. P. Jegannathan, R. Saha, dan S. K. Addya, "A time series forecasting approach to minimize cold start time in cloud-serverless platform," dalam Proc. IEEE BlackSeaCom, 2022, hlm. 325–330, doi: 10.1109/BlackSeaCom54372.2022.9858271.',
    "nguyen": 'T. N. Nguyen, "Holistic cold-start management in serverless computing cloud with deep learning for time series," Future Generation Computer Systems, vol. 153, hlm. 312–325, 2024, doi: 10.1016/j.future.2023.12.011. Pracetak: arXiv:2304.00396.',
    "icebreaker": 'R. B. Roy, T. Patel, dan D. Tiwari, "IceBreaker: Warming serverless functions better with heterogeneity," dalam Proc. ACM ASPLOS, 2022, hlm. 753–767, doi: 10.1145/3503222.3507750.',
    "femux": 'N. Nasiri, N. Munshi, S. Moser, M. Pirvu, V. Sundaresan, D. Maier, T. Premnath, N. Böwing, S. Gopalakrishnan, dan M. Shahrad, "In-production characterization of an open source serverless platform and new scaling strategies," dalam Proc. ACM EuroSys, 2026, hlm. 2054–2072, doi: 10.1145/3767295.3769377.',
    "aapa": 'G. Zhang dkk., "AAPA: An archetype-aware predictive autoscaler with uncertainty quantification for serverless workloads on Kubernetes," arXiv:2507.05653, 2025.',
    "kondrashov": 'L. Kondrashov, B. Zhou, H. Wang, dan D. Ustiugov, "The high cost of keeping warm: Characterizing overhead in serverless autoscaling policies," arXiv:2509.03104, 2025.',
    "giftpre": 'Salesforce, "GiftEvalPretrain," dataset Hugging Face Salesforce/GiftEvalPretrain (memuat alibaba_cluster_trace_2018, azure_vm_traces_2017, borg_cluster_data_2011), diakses 25 Oktober 2026.',
    "timesfm3card": 'Google Research, "timesfm-3.0-pytorch," model card Hugging Face, revisi 43046b85, diakses 25 Oktober 2026. [Daring]. Tersedia: https://huggingface.co/google/timesfm-3.0-pytorch',
    "totocard": 'Datadog, "Toto-Open-Base-1.0," model card Hugging Face, revisi 0411ceb2, diakses 25 Oktober 2026. [Daring]. Tersedia: https://huggingface.co/Datadog/Toto-Open-Base-1.0',
    "boltcard": 'Amazon, "chronos-bolt-small," model card Hugging Face, revisi 772f3d25, diakses 25 Oktober 2026. [Daring]. Tersedia: https://huggingface.co/amazon/chronos-bolt-small',
    "ttmcard": 'IBM, "granite-timeseries-ttm-r2," model card Hugging Face, diakses 25 Oktober 2026. [Daring]. Tersedia: https://huggingface.co/ibm-granite/granite-timeseries-ttm-r2',
    "little": 'J. D. C. Little, "A proof for the queuing formula: L = λW," Operations Research, vol. 9, no. 3, hlm. 383–387, 1961.',
}
ORDER: list[str] = []


def c(*keys: str) -> str:
    """Citation marker; numbers follow order of first use."""
    for k in keys:
        if k not in REFS:
            raise KeyError(k)
        if k not in ORDER:
            ORDER.append(k)
    return "[" + ", ".join(str(ORDER.index(k) + 1) for k in keys) + "]"


# ---------------------------------------------------------------- document helpers
doc = Document()
N = {"tab": 0, "fig": 0}


def setup() -> None:
    s = doc.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for lvl, size in [(1, 13), (2, 11.5), (3, 11)]:
        h = doc.styles[f"Heading {lvl}"]
        h.font.name, h.font.size, h.font.bold = "Times New Roman", Pt(size), True
        h.font.color.rgb = None
        rf = h.element.rPr.rFonts
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(a), None)
        rf.set(qn("w:eastAsia"), "Times New Roman")
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(12 if lvl == 1 else 8), Pt(4)


def runs(par, text: str, size: float | None = None) -> None:
    """Add text with **bold** and _italic_ spans."""
    for part in re.split(r"(\*\*.+?\*\*|_[^_]+?_)", text):
        if not part:
            continue
        bold, ital = part.startswith("**"), part.startswith("_") and part.endswith("_")
        r = par.add_run(part[2:-2] if bold else part[1:-1] if ital else part)
        r.bold, r.italic = bold or None, ital or None
        if size:
            r.font.size = Pt(size)


def P(text: str, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size: float | None = None, indent: bool = False):
    p = doc.add_paragraph()
    p.alignment = align
    if indent:
        p.paragraph_format.left_indent = Cm(0.6)
    runs(p, text, size)
    return p


def H(text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def bullets(items: list[str], numbered: bool = False) -> None:
    for it in items:
        p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        runs(p, it)


def border(cell, edge: str, sz: int = 8) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single"), e.set(qn("w:sz"), str(sz)), e.set(qn("w:color"), "000000")
    b.append(e)


def table(header: list[str], rows: list[list[str]], caption: str, widths_cm: list[float], size: float = 8.5,
          bold_cells: set[tuple[int, int]] = frozenset()) -> None:
    """Booktabs-style table: rules above/below header and at the bottom; caption above."""
    N["tab"] += 1
    cap = doc.add_paragraph()
    runs(cap, f"**Tabel {N['tab']}.** {caption}", 9.5)
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")  # otherwise LibreOffice/Word re-flow the column widths
    t._tbl.tblPr.append(layout)
    for j, w in enumerate(widths_cm):
        t.columns[j].width = Cm(w)
    for i, row in enumerate([header, *rows]):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Cm(widths_cm[j])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = i < len(rows)  # keep table rows on one page
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            r.font.size = Pt(size)
            r.bold = (i == 0) or ((i - 1, j) in bold_cells) or None
            if i == 0:
                border(cell, "top"), border(cell, "bottom", 4)
            if i == len(rows):
                border(cell, "bottom")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(path: Path, caption: str, width_cm: float = 15.5) -> None:
    N["fig"] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(cap, f"**Gambar {N['fig']}.** {caption}", 9.5)


# ---------------------------------------------------------------- data-driven tables
DS_COLS = ["huawei_5min", "alibaba2018_5min", "azure2019_5min", "azure2019_1min", "azure2021_5min", "azure2021_1min"]
DS_LABEL = ["Huawei 5m", "Alibaba 5m", "Azure19 5m", "Azure19 1m", "Azure21 5m", "Azure21 1m"]


def accuracy_rows() -> tuple[list[list[str]], set[tuple[int, int]]]:
    acc = pd.read_csv(TAB / "accuracy.csv", index_col=0)
    rows, bold = [], set()
    for i, (model, r) in enumerate(acc.iterrows()):
        cells = []
        for j, ds in enumerate(DS_COLS):
            v = str(r[ds])
            if v.startswith("\\textbf{"):
                v = v[len("\\textbf{"):-1]
                bold.add((i, j + 1))
            cells.append(re.sub(r"(\d)\.(\d)", r"\1,\2", v).replace(", ", "; "))  # 0.41, 0.54 -> 0,41; 0,54
        rows.append([model.replace("Always zero", "Selalu nol").replace("Toto (20 samples)", "Toto (20 sampel)"), *cells])
    return rows, bold


def dataset_rows() -> list[list[str]]:
    d = pd.read_csv(TAB / "datasets.csv", index_col=0)
    return [[lab, *(str(int(d.loc[ds, k])) for k in ["series", "smooth", "erratic", "intermittent", "lumpy"])]
            for ds, lab in zip(DS_COLS, DS_LABEL)]


# ---------------------------------------------------------------- numbering by label (cross-references cannot drift)
TAB_ORDER = ["dataset", "model", "audit", "acc", "paired", "floors", "friedman", "perclass", "sparse", "toner", "latency",
             "serverless", "vm", "ckpt"]
FIG_ORDER = ["cd", "context", "pareto", "tradeoff"]


def T(label: str) -> str:
    return f"Tabel {TAB_ORDER.index(label) + 1}"


def G(label: str) -> str:
    return f"Gambar {FIG_ORDER.index(label) + 1}"


def ltable(label: str, *args, **kw) -> None:
    assert N["tab"] == TAB_ORDER.index(label), (label, N["tab"])
    table(*args, **kw)


def lfigure(label: str, *args, **kw) -> None:
    assert N["fig"] == FIG_ORDER.index(label), (label, N["fig"])
    figure(*args, **kw)


def num(v: float, d: int = 3) -> str:
    return f"{v:.{d}f}".replace(".", ",")


def sup(n: int) -> str:
    return str(n).translate(str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹"))


def idname(name: str) -> str:
    return name.replace("Always zero", "Selalu nol").replace("Toto (20 samples)", "Toto (20 sampel)")


def friedman_rows() -> list[list[str]]:
    f = pd.read_csv(TAB / "friedman_nemenyi.csv", index_col=[0, 1])
    rows = []
    for ds, lab in zip(DS_COLS, DS_LABEL):
        r = f.loc[ds].iloc[-1]  # longest horizon
        rows.append([lab, str(int(r.n)), str(int(r.k)), num(r.chi2, 1), "10" + sup(int(round(r.log10_p))), num(r.CD, 2),
                     f"{idname(r.best)} ({num(r.rank_best, 2)})", f"{idname(r.second)} ({num(r.rank_second, 2)})"])
    return rows


def pval(p: float) -> str:
    if p >= 1e-3:
        return num(p, 3)
    mant, exp = f"{p:.1e}".split("e")
    return f"{mant.replace('.', ',')} × 10" + sup(int(exp))


def paired_rows() -> list[list[str]]:
    t = pd.read_csv(TAB / "paired_tests.csv")
    names = dict(zip(["timesfm3", "chronos2", "toto", "nbeats", "lstm", "lightgbm", "mstl_arima"],
                     ["TimesFM-3", "Chronos-2", "Toto", "N-BEATS", "LSTM", "LightGBM", "MSTL-ARIMA"]))
    return [[lab, names.get(r.tsfm, r.tsfm), names.get(r.baseline, r.baseline), str(int(r.n)),
             f"{num(r.ratio)} [{num(r.lo, 2)}; {num(r.hi, 2)}]", pval(r.p_wilcoxon), pval(r.p_holm)]
            for lab, r in zip(DS_LABEL, t.itertuples())]


def floor_rows() -> list[list[str]]:
    t = pd.read_csv(TAB / "floor_sensitivity.csv")
    label = {"floor 1e-09": "batas bawah 10⁻⁹", "floor 0.001": "batas bawah 0,001", "floor 0.01": "batas bawah 0,01 (utama)",
             "floor 0.1": "batas bawah 0,1", "ratio of sums": "rasio jumlah galat", "median ratio": "median rasio"}
    rows = []
    for metric, lab in label.items():
        g = t[t.metric == metric].set_index("dataset")
        rows.append([lab, *(f"{num(g.loc[ds, 'tsfm_value'])} / {num(g.loc[ds, 'baseline_value'])}" for ds in DS_COLS)])
    return rows


def sparse_rows() -> list[list[str]]:
    t = pd.read_csv(TAB / "azure2021_active_silent.csv")
    t = t[t.dataset == "azure2021_1min"]
    return [[r.model.replace("Always zero", "Selalu nol"), num(r.relMASE_active), num(r.rMAE_sum_active),
             num(r.overforecast_silent, 5)] for r in t.itertuples()]


# ---------------------------------------------------------------- content (v0.4: revisi review kedua, gaya humanizer-indonesia)
def front_matter() -> None:
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Apakah Model Fondasi Deret Waktu Menghasilkan Penskalaan Otomatis yang Lebih Baik? "
                  "Evaluasi Zero-Shot Berbasis Keputusan dengan Anggaran CPU pada Trace Cloud dan Serverless Publik")
    r.bold, r.font.size = True, Pt(15)
    P("[Nama Penulis]¹, [Nama Penulis]²", WD_ALIGN_PARAGRAPH.CENTER)
    P("¹[Afiliasi], ²[Afiliasi], [email korespondensi]", WD_ALIGN_PARAGRAPH.CENTER, 10)
    P("_DRAFT v0.5, 26 Oktober 2026. Angka diambil dari results/ dan paper/tables/._", WD_ALIGN_PARAGRAPH.CENTER, 9)

    H("Abstrak", 1)
    P("Model fondasi deret waktu (_time-series foundation model_, TSFM) dapat meramalkan beban kerja tanpa harus dilatih ulang "
      "untuk setiap layanan. Kemampuan ini menarik untuk penskalaan otomatis (_autoscaling_) pada platform _serverless_ dan "
      "_middleware_, tetapi satu-satunya evaluasi yang kami temukan pada rekaman jejak (_trace_) beban serverless, yaitu pada data "
      "Huawei, justru mendapati TSFM kalah dari model _seasonal naive_. Pada penelitian ini kami mengevaluasi delapan TSFM rilis "
      "2024 hingga 2026 secara _zero-shot_ terhadap dua belas model pembanding statistik dan terlatih. Evaluasi tersebut memakai "
      "empat trace publik, yaitu Azure Functions 2019 dan 2021, Alibaba 2018, dan Huawei, dengan total 920 deret. Ada tiga sisi "
      "yang diukur: akurasi pada horizon 5 hingga 60 menit, biaya inferensi di CPU, serta kualitas keputusan penskalaan dalam "
      "simulasi berbasis trace. Pada keenam kombinasi dataset dan resolusi, TSFM terbaik (TimesFM-3, Chronos-2, atau Toto, dipilih "
      "pada periode validasi) menurunkan galat relatif 7 hingga 26% dibanding model pembanding terbaik, dan seluruh selisih itu "
      "signifikan menurut uji Wilcoxon berpasangan dengan koreksi Holm (p ≤ 0,022). Arah hasilnya bertahan pada berbagai cara "
      "menangani galat nol, meskipun pada Azure 2021 selisih menurut median rasio hampir nol, dan korpus pelatihan awal sebagian besar "
      "TSFM kemungkinan besar memuat trace Alibaba. Ketika "
      "konfigurasi eksperimen studi sebelumnya kami ulangi dengan model yang lebih baru, hasil negatifnya tidak muncul lagi. "
      "Keunggulan akurasi itu tidak menghasilkan perbaikan keputusan yang sama besar di setiap trace. Untuk pemanasan awal instance "
      "pada Azure 2019, rencana kapasitas berbasis Chronos-2 mengurangi proksi instance menganggur sekitar 69% dibanding _keep-alive_ "
      "reaktif dan sekitar 48% dibanding LightGBM pada proksi _cold start_ 5%. Pada aplikasi Azure 2021 yang sangat jarang "
      "dipanggil, manfaatnya tidak konsisten, dan pada mesin virtual Alibaba yang sudah memakai autoscaler reaktif TSFM hanya "
      "setara dengan MSTL-ARIMA. Dari sisi biaya, Chronos-2 dan Toto menawarkan kompromi akurasi dan latensi yang menarik pada 0,7 "
      "hingga 1,7 detik untuk 64 deret di CPU laptop, sedangkan TimesFM-3 memiliki galat agregat terendah dengan latensi 3,3 detik.")
    P("**Kata kunci:** penskalaan otomatis, serverless, peramalan beban kerja, model fondasi deret waktu, zero-shot, cold start.")

    H("Abstract", 1)
    P("_Time-series foundation models (TSFMs) can forecast workloads without being retrained for every service, which makes "
      "them attractive for autoscaling serverless and middleware platforms. The only evaluation we found on serverless traces, on "
      "Huawei data, reported that TSFMs lose to a seasonal naive forecaster. We evaluate eight TSFMs released between 2024 and 2026 "
      "in a zero-shot setting against twelve statistical and trained baselines on four public traces (Azure Functions 2019 and "
      "2021, Alibaba 2018, and Huawei; 920 series), measuring forecast accuracy at 5 to 60 minute horizons, CPU inference cost, "
      "and the quality of autoscaling decisions in a trace-driven simulation. On all six dataset-resolution combinations, the best "
      "TSFM (TimesFM-3, Chronos-2, or Toto, selected on validation data) reduces the relative error of the best baseline by 7 to 26%, "
      "and every difference is significant under a paired Wilcoxon test with Holm correction (p <= 0.022). The direction holds "
      "under several treatments of zero errors, although on Azure 2021 the median ratio is nearly tied, and most TSFMs likely saw "
      "the Alibaba trace during pretraining. Re-running the earlier "
      "study's setting with newer models reverses its negative result. Accuracy gains do not translate into equal decision gains "
      "on every trace. For pre-warming on Azure 2019, Chronos-2-based capacity plans cut a proxy of idle instances by about 69% relative "
      "to reactive keep-alive and about 48% relative to LightGBM at a 5% cold-start proxy; on the very sparse Azure 2021 "
      "applications the benefit is inconsistent, and for Alibaba virtual machines already governed by a reactive autoscaler, TSFMs "
      "only match MSTL-ARIMA. Chronos-2 and Toto offer an attractive accuracy-latency trade-off at 0.7 to 1.7 seconds per 64 series "
      "on a laptop CPU, while TimesFM-3 has the lowest aggregate error at 3.3 seconds._")
    P("_**Keywords:** autoscaling, serverless, workload forecasting, time-series foundation model, zero-shot, cold start._")


def intro() -> None:
    H("1. Pendahuluan", 1)
    P("Autoscaler reaktif, misalnya Horizontal Pod Autoscaler (HPA) pada Kubernetes " + c("hpa") + ", baru menambah kapasitas "
      "setelah utilisasi naik, sehingga lonjakan beban telanjur dilayani dengan kapasitas yang kurang. Pada platform _serverless_ "
      "(nirserver), keterlambatan ini muncul sebagai _cold start_: permintaan ke fungsi yang tidak memiliki instance hangat harus "
      "menunggu instance baru disiapkan. Shahrad dkk. " + c("shahrad") + " menemukan bahwa sebagian besar aplikasi Azure Functions "
      "jarang dipanggil, sehingga kebijakan _keep-alive_ dengan durasi tetap selalu berhadapan dengan pilihan antara _cold start_ "
      "yang sering dan memori yang terbuang. Penskalaan prediktif berupaya keluar dari pilihan itu dengan meramalkan beban lebih "
      "dahulu " + c("roy") + ". Kendalanya, model peramal seperti ARIMA (_autoregressive integrated moving average_) atau LSTM "
      "(_long short-term memory_) umumnya dilatih untuk setiap layanan atau klaster. Bagi operator yang mengelola ribuan fungsi, "
      "pekerjaan melatih dan memantau ribuan model tersebut menjadi biaya operasional tersendiri.")
    P("Model fondasi deret waktu (TSFM) menawarkan cara lain. Satu model dilatih awal pada korpus deret waktu yang besar, kemudian "
      "dipakai langsung pada deret baru tanpa pelatihan tambahan (_zero-shot_) " + c("chronos", "timesfm", "moirai", "ttm") +
      ". Sejauh penelusuran kami, evaluasi langsung pada beban serverless baru dilakukan oleh Toner dkk. " + c("toner") + ", dan "
      "hasilnya negatif: pada empat subset trace Huawei, model _seasonal naive_ yang hanya mengulang nilai pada jam yang sama sehari "
      "sebelumnya mengalahkan semua TSFM yang diuji. Studi tersebut memiliki empat keterbatasan. Datanya berasal dari satu penyedia, "
      "dan _checkpoint_ yang diuji berasal dari generasi awal dengan ukuran kecil (Chronos-tiny dan Moirai-small). Evaluasinya juga "
      "hanya mengukur akurasi tanpa menguji keputusan penskalaan, sementara biaya inferensi di CPU hanya dilaporkan untuk sebagian "
      "model.")
    P("Penelitian ini berupaya menjawab tiga pertanyaan. **RQ1:** apakah TSFM generasi terbaru dapat menyamai atau mengungguli "
      "model pembanding statistik dan terlatih pada beban cloud dan serverless, dan pada pola beban seperti apa? **RQ2:** berapa "
      "biaya inferensinya di CPU, dan model mana yang dapat dijalankan di dalam siklus kontrol autoscaler? **RQ3:** apakah keunggulan "
      "akurasi tersebut menghasilkan keputusan penskalaan yang lebih baik?")
    P("Kontribusi penelitian ini adalah sebagai berikut:")
    bullets([
        "Evaluasi tolok ukur (_benchmark_) _zero-shot_ untuk delapan TSFM terhadap dua belas model pembanding pada empat trace "
        "publik dan enam kombinasi dataset dan resolusi. Totalnya 920 deret dan 150.120 pasangan titik awal dan deret untuk setiap "
        "model, dilengkapi interval kepercayaan _bootstrap_ serta uji Friedman dan Nemenyi.",
        "Evaluasi ulang dengan konfigurasi eksperimen Toner dkk. " + c("toner") + ". Hasilnya mengindikasikan bahwa hasil negatif "
        "mereka tidak berlaku untuk TSFM generasi 2025 hingga 2026, tetapi tetap berlaku untuk model yang berukuran kecil.",
        "Evaluasi berbasis keputusan melalui simulator penskalaan untuk serverless dan mesin virtual (_virtual machine_, VM), "
        "lengkap dengan pseudocode dan definisi metrik. Simulator ini membandingkan kurva kompromi (_trade-off_) setiap kebijakan "
        "pada tingkat risiko yang sama, termasuk adaptasi kebijakan histogram hibrida Shahrad dkk. " + c("shahrad") + ".",
        "Profil biaya CPU untuk seluruh model, ablasi panjang konteks yang menunjukkan bahwa konteks optimal bergantung pada "
        "resolusi data, serta audit kebocoran data pelatihan awal untuk setiap TSFM.",
    ], numbered=True)
    P("Pada tingkat akurasi, galat TSFM terbaik 7 hingga 26% lebih kecil daripada model pembanding terbaik pada keenam kombinasi "
      "dataset dan resolusi, termasuk aplikasi serverless yang sangat jarang dipanggil, dan seluruh selisih tersebut signifikan. Pada tingkat keputusan, keunggulan tersebut "
      "menghasilkan penghematan yang besar untuk pemanasan awal (_pre-warming_) instance pada Azure 2019, tidak konsisten pada "
      "Azure 2021, dan hampir hilang untuk VM yang sudah memakai kontrol reaktif.")


def related() -> None:
    H("2. Kajian Terkait", 1)
    H("2.1 Peramalan beban kerja untuk penskalaan otomatis", 2)
    P("Autoscaler yang dipakai di lingkungan produksi umumnya bersifat reaktif. HPA, misalnya, menghitung jumlah replika dari "
      "utilisasi yang teramati dan memakai jendela stabilisasi sebelum menurunkan kapasitas " + c("hpa") + ". Roy dkk. " + c("roy") +
      " memakai model ARMA untuk meramalkan beban dan menyiapkan VM lebih awal. Untuk serverless, Shahrad dkk. " + c("shahrad") +
      " mengusulkan kebijakan histogram hibrida yang menentukan jendela pemanasan awal dan _keep-alive_ dari distribusi waktu "
      "antarinvokasi setiap aplikasi. Karakterisasi trace Azure " + c("shahrad", "zhang21") + ", Huawei " + c("joosen") + ", dan "
      "Alibaba " + c("guo") + " memperlihatkan dua sifat yang menyulitkan peramalan. Volume permintaan antarfungsi dapat berbeda "
      "hingga sembilan orde besaran, dan banyak fungsi hanya aktif sesekali.")
    P("Peramalan juga sudah lama dipakai untuk menekan _cold start_. Golec dkk. " + c("golec") + " merangkum berbagai pendekatan "
      "itu dalam kajian sistematis. Jegannathan dkk. " + c("jegannathan") + " memakai SARIMA untuk memperkirakan kedatangan "
      "permintaan dan menyesuaikan jumlah kontainer, sedangkan Nguyen " + c("nguyen") + " memakai _deep learning_ untuk deret waktu "
      "dengan tujuan yang sama. IceBreaker " + c("icebreaker") + " memperkirakan peluang invokasi berikutnya untuk memilih jenis "
      "node yang dipanaskan, dan melaporkan penurunan biaya _keep-alive_ sebesar 45%. FeMux " + c("femux") + " menggabungkan "
      "beberapa model peramal ringan pada Knative dan melaporkan pengurangan _cold start_ hingga 64% dibanding FaasCache, "
      "IceBreaker, dan Aquatope. AAPA " + c("aapa") + " mengelompokkan jendela beban Azure Functions ke dalam beberapa arketipe "
      "sebelum memilih strategi penskalaan. Kondrashov dkk. " + c("kondrashov") + " mengukur bahwa pergantian instance akibat "
      "kebijakan penskalaan menambah beban CPU sebesar 10 sampai 40% dari siklus yang dipakai untuk menangani permintaan, dan "
      "memori yang dialokasikan 2 sampai 10 kali lebih besar daripada yang benar-benar dipakai. Seluruh studi tersebut memakai "
      "model peramal yang dilatih untuk platform atau aplikasi tertentu. Sejauh penelusuran kami, belum ada yang menguji TSFM "
      "_zero-shot_ untuk keperluan ini.")
    H("2.2 Model fondasi deret waktu", 2)
    P("Chronos " + c("chronos") + " mengubah nilai deret menjadi token, lalu melatih model bahasa di atas token tersebut. "
      "Chronos-Bolt " + c("boltcard") + " mempercepat proses itu dengan memakai _patch_ dan keluaran kuantil langsung, sedangkan "
      "Chronos-2 " + c("chronos2") + " menambahkan dukungan untuk deret multivariat dan kovariat. TimesFM " + c("timesfm") + " adalah "
      "_transformer_ yang hanya memakai dekoder, dan versi 2.5 " + c("timesfm25card") + " serta 3 " + c("timesfm3", "timesfm3card") +
      " menambahkan keluaran kuantil dan konteks hingga 16 ribu langkah. Moirai " + c("moirai") + " dilatih pada korpus LOTSA, "
      "kemudian Moirai 2.0 " + c("moirai2") + " beralih ke arsitektur dekoder yang lebih kecil. TTM " + c("ttm", "ttmcard") +
      " memakai arsitektur MLP-Mixer dengan ukuran di bawah lima juta parameter, TiRex " + c("tirex") + " dibangun di atas xLSTM, "
      "dan Toto " + c("toto", "totocard") + " dilatih dengan porsi besar metrik _observability_ milik Datadog. Klaim keunggulan "
      "model-model ini sebagian besar berasal dari tolok ukur umum seperti GIFT-Eval " + c("gifteval") + ", yang hanya memuat "
      "sedikit data beban kerja cloud.")
    H("2.3 TSFM pada data cloud", 2)
    P("Woo dkk. " + c("cloudops") + " melatih model awal pada data CloudOps yang mencakup Alibaba 2018, Azure VM 2017, dan Borg "
      "2011, sehingga dataset tersebut kemungkinan sudah masuk ke korpus pelatihan awal model turunannya. Cohen dkk. " + c("toto") +
      " membangun tolok ukur _observability_ BOOM dari data internal Datadog. Evaluasi _zero-shot_ yang paling dekat dengan "
      "penelitian ini adalah studi Toner dkk. " + c("toner") + ". Sejauh penelusuran kami, belum ada studi yang menguji TSFM sampai "
      "ke tahap keputusan penskalaan, dan belum ada pula yang mengukur biaya CPU seluruh model dalam kondisi yang sama.")


def method() -> None:
    H("3. Metodologi", 1)
    H("3.1 Dataset", 2)
    P("Penelitian ini memakai empat _trace_ publik. **Azure Functions 2019** " + c("shahrad") + " berisi jumlah invokasi per menit "
      "untuk setiap fungsi selama 14 hari. Jumlah tersebut kami agregasikan ke tingkat aplikasi, lalu kami pertahankan 13.057 "
      "aplikasi yang muncul pada ke-14 hari. Dari kumpulan itu kami mengambil 50 aplikasi dengan volume tertinggi, ditambah sampel "
      "acak hingga 125 aplikasi untuk setiap kelas pola permintaan menurut Syntetos dan Boylan " + c("syntetos") + ", sehingga "
      "diperoleh 550 aplikasi. **Azure Functions 2021** " + c("zhang21") + " berisi catatan untuk setiap invokasi. Waktu "
      "kedatangannya dihitung sebagai waktu selesai dikurangi durasi, kemudian dikelompokkan per menit untuk 119 aplikasi.")
    P("**Alibaba 2018** " + c("guo") + " menyediakan utilisasi CPU per mesin dengan interval pencuplikan yang tidak teratur "
      "(median 60 detik, maksimum sekitar 640 detik), sehingga kami memakai bin 5 menit. Satu gangguan pencatatan yang mengenai "
      "seluruh mesin selama 5,4 jam pada hari kedua kami isi dengan nilai pada jam yang sama sehari sebelumnya. Kami kemudian "
      "mengambil 200 mesin secara acak dari 3.787 mesin yang lolos filter cakupan data 90%, ditambah satu deret rata-rata klaster. "
      "**Huawei** " + c("joosen") + " berupa 50 kanal hasil kurasi Toner dkk. " + c("toner") + " dengan resolusi 5 menit selama "
      "30 hari.")
    P("Data Azure juga dievaluasi pada resolusi 5 menit dengan menjumlahkan invokasi setiap 5 menit. Seluruh stempel waktu pada "
      "trace sudah dianonimkan, sehingga tidak ada model yang memakai fitur kalender. Jumlah deret dan kelas polanya dirangkum pada "
      + T("dataset") + ".")
    ltable("dataset", ["Dataset", "Deret", "Smooth", "Erratic", "Intermittent", "Lumpy"], dataset_rows(),
           "Jumlah deret untuk setiap kombinasi dataset dan resolusi, beserta kelas pola Syntetos dan Boylan yang dihitung pada "
           "data latih.", [3.6, 1.8, 1.8, 1.8, 2.2, 1.8])

    H("3.2 Model", 2)
    P(T("model") + " memuat delapan TSFM yang dievaluasi. Revisi _checkpoint_ (bobot model) Hugging Face yang kami pakai "
      "dicantumkan pada Lampiran A. Seluruh model dijalankan secara _zero-shot_ di CPU dengan panjang konteks C = 2048 langkah, "
      "kecuali Alibaba yang memakai 1440 langkah karena riwayat datanya lebih pendek. Toto dijalankan dalam mode rata-rata "
      "analitik, dan varian dengan 20 sampel dilaporkan terpisah. Moirai 2.0 tidak memiliki keluaran rata-rata, sehingga kami "
      "memakai mediannya sebagai ramalan titik.")
    ltable("model", ["Model", "Checkpoint", "Ukuran", "Keluaran"], [
        ["TimesFM-3", "google/timesfm-3.0-pytorch", "330M", "kuantil"],
        ["Chronos-2", "amazon/chronos-2", "120M", "kuantil"],
        ["Toto", "Datadog/Toto-Open-Base-1.0", "151M", "rata-rata / sampel"],
        ["TimesFM-2.5", "google/timesfm-2.5-200m-pytorch", "200M", "kuantil"],
        ["TiRex", "NX-AI/TiRex", "35M", "kuantil"],
        ["Moirai-2-S", "Salesforce/moirai-2.0-R-small", "kecil", "kuantil"],
        ["Chronos-Bolt-S", "amazon/chronos-bolt-small", "48M", "kuantil"],
        ["TTM-r2", "ibm-granite/granite-timeseries-ttm-r2", "< 5M", "titik"],
    ], "TSFM yang dievaluasi. Revisi _checkpoint_ tercantum pada Lampiran A.", [2.8, 6.4, 1.6, 2.4])
    P("Model pembanding terbagi atas tiga kelompok. Kelompok pertama adalah model sederhana, yaitu model yang selalu meramalkan "
      "nol, _naive_ (nilai terakhir), _seasonal naive_ (nilai pada jam yang sama sehari sebelumnya), dan SWA-3 (rata-rata nilai "
      "pada jam yang sama selama tiga hari terakhir). Kelompok kedua adalah model statistik lokal yang dijalankan dengan pustaka "
      "StatsForecast " + c("statsforecast") + ", yaitu ETS (_exponential smoothing_) tanpa komponen musiman, MSTL " + c("mstl") +
      " dengan ETS atau ARIMA untuk komponen tren, Croston " + c("croston") + ", dan TSB " + c("tsb") + ". MSTL hanya dijalankan "
      "pada resolusi 5 menit, karena satu kali penaksiran parameternya pada data 1 menit memakan waktu 4 hingga 9 detik. Kelompok "
      "ketiga adalah model terlatih global yang dilatih untuk setiap dataset: LightGBM " + c("lightgbm") + " rekursif dengan fitur "
      "_lag_ yang dinormalisasi, serta N-BEATS " + c("nbeats") + " dan LSTM " + c("lstm") + " yang dijalankan dengan "
      "NeuralForecast " + c("neuralforecast") + " dan dihentikan lebih awal (_early stopping_) berdasarkan data validasi. Model "
      "statistik melihat riwayat 7 hari, lebih panjang daripada konteks TSFM, sehingga perbandingan ini cenderung menguntungkan "
      "model pembanding.")

    H("3.3 Protokol evaluasi", 2)
    P("Setiap dataset dibagi secara kronologis menjadi data latih, validasi, dan uji, dengan porsi 10/2/2 hari untuk Azure, 5/1/2 "
      "hari untuk Alibaba, dan 24/2/4 hari untuk Huawei. Ramalan dibuat dari titik awal (_forecast origin_) setiap jam sepanjang "
      "periode validasi dan uji. Horizon yang dievaluasi adalah 5, 15, dan 60 menit, ditambah 1 menit untuk data beresolusi 1 menit. "
      "Ramalan pada periode validasi dipakai untuk membentuk kuantil konformal " + c("vovk") + " bagi model yang hanya menghasilkan "
      "ramalan titik. Residu validasi itu dinormalisasi dengan skala setiap deret, lalu dikumpulkan untuk setiap langkah horizon. "
      "Ramalan yang bernilai negatif dipotong menjadi nol.")
    P("Untuk deret i dan horizon H, MAE dihitung dari seluruh langkah 1 sampai H di semua titik awal pada periode uji, sehingga "
      "“horizon 60 menit” pada tabel berarti rata-rata galat untuk ramalan 1 sampai 60 menit ke depan. Metrik utama kami "
      "adalah relMASE = exp(rata-rata_i log(max(MAE_i / MAE_i^SN, 0,01))), dengan MAE_i^SN galat _seasonal naive_ pada deret yang "
      "sama. Rasio ini setara dengan rasio MASE (_mean absolute scaled error_) " + c("hyndman") + ", dan nilai di bawah 1 berarti "
      "lebih baik daripada _seasonal naive_. Batas bawah 0,01 mencegah satu deret dengan galat nol menarik rata-rata geometrik ke "
      "arah nol tanpa batas. Masalah ini nyata pada data yang sangat jarang: pada 8 dari 77 deret Azure 2021 yang dapat dihitung, "
      "LSTM, LightGBM, dan _naive_ memiliki MAE tepat nol. Deret dengan MAE_i^SN = 0 dikeluarkan karena rasionya tidak terdefinisi. "
      "Sebagai metrik pendamping yang tidak memerlukan batas bawah, kami juga melaporkan rasio jumlah galat Σ_i MAE_i / Σ_i MAE_i^SN.")
    P("Interval kepercayaan 95% dihitung dengan 1000 resampel _bootstrap_ atas deret. Interval ini mengukur ketidakpastian antarderet "
      "untuk periode uji yang tetap, dan tidak mencakup variasi antarperiode. Untuk uji signifikansi kami memakai uji Friedman "
      "dengan koreksi untuk nilai seri (_ties_), diikuti uji lanjut Nemenyi " + c("demsar") + " pada peringkat MAE setiap deret. "
      "Jumlah deret, jumlah model, statistik χ², dan nilai p setiap kombinasi dilaporkan apa adanya pada " + T("friedman") + ".")
    P("Klaim bahwa satu model lebih baik daripada model lain kami uji secara berpasangan, tidak dari tumpang tindih interval "
      "kepercayaan. Untuk setiap kombinasi, TSFM terbaik dan model pembanding terbaik dipilih berdasarkan relMASE pada **periode "
      "validasi**, sehingga pemilihan pasangan tidak melihat data uji. Selisih log rasio kedua model pada setiap deret di periode uji "
      "kemudian diuji dengan uji peringkat bertanda Wilcoxon, dan interval kepercayaan rasio geometriknya dihitung dengan "
      "_bootstrap_ berpasangan. Nilai p dikoreksi dengan metode Holm untuk enam pengujian. Karena hasil pada data yang sangat jarang "
      "peka terhadap perlakuan galat nol, kami juga melaporkan perbandingan yang sama dengan batas bawah 10⁻⁹, 0,001, dan 0,1, rasio "
      "jumlah galat, dan median rasio (" + T("floors") + ").")

    H("3.4 Biaya komputasi", 2)
    P("Latensi diukur secara terpisah pada Apple M3 Pro (11 inti, RAM 18 GB) dalam kondisi tanpa beban lain. Pengukuran memakai "
      "4 utas CPU, horizon 60, ukuran _batch_ 1 dan 64, serta konteks 512 dan 2048. Setiap konfigurasi diukur 30 kali setelah 5 kali "
      "pemanasan, dan yang kami laporkan adalah mediannya (p50). Waktu pelatihan model pembanding dicatat dari eksekusi penuh. "
      "Latensi inferensi model pembanding terlatih tidak kami ukur dengan protokol yang sama.")

    H("3.5 Simulasi penskalaan otomatis", 2)
    P("Simulator bekerja pada agregat per menit dan tidak mensimulasikan permintaan individual. Karena itu metrik risikonya kami sebut "
      "**proksi** _cold start_. Untuk setiap aplikasi i dan menit t, langkahnya sebagai berikut:")
    bullets([
        "Kebutuhan instance dihitung dengan hukum Little " + c("little") + ": N_it = ⌈λ_it · d_i / 60⌉, dengan λ_it jumlah invokasi "
        "pada menit t dan d_i durasi rata-rata aplikasi dalam detik. Durasi d_i dihitung dari data latih saja, sehingga 2 aplikasi "
        "Azure 2019 dan 16 aplikasi Azure 2021 yang tidak aktif selama periode latih dikeluarkan dari simulasi.",
        "Setiap kebijakan menentukan kapasitas hangat K_it untuk menit tersebut.",
        "Kekurangan instance dihitung sebagai max(0, N_it − K_it), dan kelebihan sebagai max(0, K_it − N_it).",
        "Proksi rasio _cold start_ = Σ_it max(0, N_it − K_it) / Σ_it N_it, dan instance menganggur per kebutuhan = "
        "Σ_it max(0, K_it − N_it) / Σ_it N_it. Pembilang dan penyebut keduanya dijumlahkan atas seluruh aplikasi sebelum dibagi.",
    ], numbered=True)
    P("Sebagai contoh, sebuah aplikasi dengan durasi 4 detik menerima 30 invokasi pada menit t, sehingga N_t = ⌈30 · 4/60⌉ = 2. "
      "Bila kebijakan hanya menyiapkan K_t = 1 instance, kekurangannya 1 dan menit ini menyumbang 1 dari 2 unit kebutuhan ke proksi "
      "_cold start_. Bila pada menit berikutnya tidak ada invokasi tetapi K = 1, menit itu menyumbang 1 instance-menit menganggur. "
      "Simulator tidak melacak instance yang sedang sibuk melewati batas menit dan tidak memodelkan waktu _cold start_ itu sendiri, "
      "sehingga angka yang dihasilkan adalah estimasi relatif antarkebijakan.")
    P("Kami membandingkan empat kebijakan serverless. Kebijakan reaktif memakai _keep-alive_ dengan durasi tetap K menit, "
      "K ∈ {1, 2, 5, 10, 20, 30, 60}. Kebijakan prediktif menyiapkan kapasitas untuk 60 menit berikutnya pada setiap titik awal per "
      "jam: K_it = ⌈g(q̂_τ) · d_i / 60⌉, dengan q̂_τ kuantil τ dari ramalan dan g(x) = x bila x ≥ θ dan 0 bila sebaliknya. Ambang "
      "aktivasi θ mencegah ramalan positif yang sangat kecil langsung dibulatkan menjadi satu instance. Pasangan (τ, θ) dengan "
      "τ ∈ {0,5; …; 0,9} dan θ ∈ {0; 0,01; 0,05; 0,1; 0,25; 0,5; 1} invokasi per menit dipilih pada **periode validasi**: semua "
      "pasangan yang Pareto-optimal di validasi kami pertahankan, lalu dievaluasi pada periode uji. Untuk model titik, kuantil "
      "konformal dikalibrasi pada periode validasi yang sama, sehingga pemilihan ini sedikit optimistis bagi model tersebut. "
      "Kebijakan hibrida mengambil nilai maksimum dari kebijakan prediktif dan _keep-alive_ 10 menit. Kebijakan keempat adalah "
      "**adaptasi** kebijakan histogram hibrida Shahrad dkk. " + c("shahrad") + " tanpa komponen ARIMA dan tanpa pembaruan daring: "
      "histogram waktu menganggur setiap aplikasi (rentang 4 jam) dipelajari satu kali dari data sebelum periode uji, jendela "
      "pemanasan awal ditetapkan dari persentil ke-5, dan jendela _keep-alive_ dari persentil ke-99, masing-masing dengan margin "
      "10%. Aplikasi dengan histogram yang tidak representatif, termasuk aplikasi yang di versi aslinya ditangani ARIMA, kembali ke "
      "_keep-alive_ 4 jam.")
    P("**VM (Alibaba, 5 menit).** Kebutuhan replika dihitung sebagai R = ⌈u/100 · 20/0,7⌉ untuk layanan berkapasitas 20 replika "
      "pada utilisasi 100% dengan target utilisasi 70%. Kebijakan reaktif meniru HPA, yaitu memakai permintaan satu langkah "
      "sebelumnya dengan stabilisasi penurunan satu langkah, dan target utilisasi penyediaan ∈ {0,4; …; 0,8} menjadi parameter "
      "pengaturnya. Kebijakan prediktif dan hibrida memakai kuantil τ seperti pada serverless tanpa ambang aktivasi, karena "
      "utilisasi bernilai kontinu. Risiko diukur sebagai fraksi interval yang kekurangan kapasitas, dan biaya diukur sebagai "
      "replika berlebih per replika yang dibutuhkan.")
    P("Setiap kebijakan memiliki parameter pengatur, sehingga yang kami bandingkan adalah kurva kompromi. Biaya setiap kebijakan "
      "diinterpolasi pada tingkat risiko yang sama.")

    H("3.6 Audit kebocoran data pelatihan awal", 2)
    P("Evaluasi _zero-shot_ hanya sah bila data uji tidak ikut dipakai saat pelatihan awal. Kami memeriksa _model card_ "
      "(dokumentasi resmi model) dan paper setiap TSFM (" + T("audit") + "). Korpus GiftEvalPretrain " + c("giftpre") + ", yang "
      "merupakan subset dari LOTSA " + c("moirai") + ", memuat trace Alibaba 2018 yang telah diolah Woo dkk. " + c("cloudops") +
      ". Kami membedakan tiga tingkat tumpang tindih: (1) trace sumber yang sama, (2) klaster dan periode yang sama tetapi tabel "
      "atau sinyal berbeda, dan (3) observasi target uji yang sama. Berdasarkan jumlah observasinya, olahan Woo dkk. kemungkinan "
      "besar berasal dari tabel penggunaan kontainer, sedangkan kami memakai tabel penggunaan mesin. Tumpang tindih pada Alibaba "
      "dengan demikian berada pada tingkat (2), sedangkan tingkat (3) tidak dapat kami pastikan maupun kami singkirkan. Karena itu "
      "hasil Alibaba untuk lima TSFM tidak kami perlakukan sebagai _zero-shot_ murni. Trace Azure Functions dan Huawei tidak "
      "ditemukan dalam dokumentasi yang kami periksa, yang belum berarti keduanya dipastikan tidak dipakai.")
    ltable("audit", ["Model", "Korpus pelatihan awal", "Sumber", "Alibaba 2018", "Azure, Huawei"], [
        ["TimesFM-2.5", "GiftEvalPretrain, Wikipedia, Google Trends, sintetis", c("timesfm25card"), "ya", "tidak disebut"],
        ["TimesFM-3", "GiftEvalPretrain (tanpa irisan fev-bench), Wikipedia, Google Trends, sintetis", c("timesfm3card"),
         "kemungkinan besar", "tidak disebut"],
        ["Chronos-2", "subset Chronos dan GiftEvalPretrain, sintetis; Alibaba disebut pada tabel data", c("chronos2"), "ya",
         "tidak disebut"],
        ["Toto", "observability Datadog, GiftEvalPretrain, subset Chronos, sintetis", c("totocard"), "kemungkinan besar",
         "tidak disebut"],
        ["Moirai-2-S", "GiftEvalPretrain dan data lain, termasuk data internal", c("moirai2"), "kemungkinan besar", "tidak disebut"],
        ["TiRex", "data Chronos, sintetis, subset GiftEvalPretrain (Azure VM, Borg disebut)", c("tirex"), "tidak disebut",
         "tidak disebut"],
        ["Chronos-Bolt-S", "korpus Chronos (tanpa trace cloud)", c("boltcard"), "tidak disebut", "tidak disebut"],
        ["TTM-r2", "data publik tanpa rincian", c("ttmcard"), "tidak diketahui", "tidak diketahui"],
    ], "Audit kebocoran data pelatihan awal berdasarkan _model card_ dan paper setiap model (revisi dokumen per 25 Oktober 2026). "
       "“Tidak disebut” berarti tidak ditemukan dalam dokumentasi, bukan dipastikan tidak dipakai.",
       [2.4, 6.6, 1.4, 2.3, 2.3], size=7.5)


def results_accuracy() -> None:
    H("4. Hasil", 1)
    H("4.1 Akurasi (RQ1)", 2)
    rows, bold = accuracy_rows()
    ltable("acc", ["Model", *DS_LABEL], rows,
           "relMASE untuk horizon hingga 60 menit, yaitu rata-rata galat langkah 1 sampai 60 menit (rata-rata geometrik dengan batas "
           "bawah rasio 0,01 dan interval kepercayaan _bootstrap_ 95%). Nilai lebih kecil lebih baik, dan _seasonal naive_ bernilai "
           "1. Huruf tebal menandai nilai terbaik pada setiap kolom, dan “–” menandai model yang tidak dijalankan. Kolom "
           "Alibaba untuk TimesFM-3, TimesFM-2.5, Chronos-2, Toto, dan Moirai-2-S bukan hasil _zero-shot_ murni (" + T("audit") + ").",
           [2.6, 2.25, 2.25, 2.25, 2.25, 2.25, 2.25], size=7, bold_cells=bold)
    P(T("acc") + " menunjukkan bahwa posisi teratas pada keenam kombinasi dataset dan resolusi selalu ditempati TimesFM-3, "
      "Chronos-2, atau Toto, tetapi besar selisihnya berbeda antardataset. Pada Azure 2019 resolusi 5 menit, relMASE TimesFM-3 "
      "sebesar 0,469 [0,43; 0,51], sedangkan model pembanding terbaik, yaitu N-BEATS, sebesar 0,586 [0,54; 0,64], atau penurunan "
      "galat sebesar 20%. Pada resolusi 1 menit selisihnya 24% "
      "(Toto 0,556 [0,51; 0,61] dan N-BEATS 0,734 [0,66; 0,81]). Pada Huawei selisihnya 11% (Chronos-2 0,467 dan N-BEATS 0,525), dan "
      "pada Alibaba 7% (TimesFM-3 0,610 dan MSTL-ARIMA 0,653). Pada horizon 5 menit, nilai TimesFM-3 dan MSTL-ARIMA di Alibaba hampir "
      "sama (0,540 dan 0,547).")
    P("Pada Azure 2021, Toto memperoleh 0,313 (5 menit) dan 0,330 (1 menit), sedangkan pembanding terbaik, yaitu LSTM, memperoleh "
      "0,364 dan 0,439, atau selisih 14% dan 25%. Model yang selalu meramalkan nol hanya kompetitif pada trace ini (0,496 dan 0,488), sedangkan pada keempat "
      "kombinasi data padat galatnya 3,9 sampai 5,4 kali galat _seasonal naive_.")
    P("Ukuran model juga berpengaruh. Model kecil, yaitu TTM-r2 dan Chronos-Bolt-small, kalah dari LightGBM pada keempat kombinasi "
      "data Azure, dan pada Azure 2019 resolusi 1 menit TTM-r2 bahkan 54% lebih buruk daripada _seasonal naive_. Dengan kata lain, "
      "keunggulan di atas hanya berlaku untuk model besar yang terbaru, dan label _zero-shot_ saja tidak menjamin akurasi.")
    ltable("paired", ["Dataset", "TSFM", "Pembanding", "n", "Rasio [CI 95%]", "p Wilcoxon", "p Holm"], paired_rows(),
           "Uji berpasangan antara TSFM terbaik dan model pembanding terbaik, keduanya dipilih pada periode validasi. Rasio adalah "
           "rata-rata geometrik MAE TSFM terhadap MAE pembanding per deret pada periode uji (lebih kecil dari 1 berarti TSFM lebih "
           "baik), dengan interval kepercayaan _bootstrap_ berpasangan.", [2.0, 1.8, 2.4, 0.9, 2.9, 2.2, 2.2], size=7.5)
    P(T("paired") + " menguji selisih tersebut secara berpasangan. Pada keenam kombinasi, rasio galat TSFM terhadap pembanding "
      "berada di bawah 1 dengan interval kepercayaan yang tidak mencakup 1, dan selisihnya signifikan setelah koreksi Holm. "
      "Selisih paling lemah terdapat pada Azure 2021 resolusi 5 menit (TimesFM-3 terhadap LSTM, rasio 0,870 [0,76; 0,96], p Holm "
      "0,022). Tumpang tindih interval pada " + T("acc") + " tidak cukup untuk menilai signifikansi karena kedua model dievaluasi "
      "pada deret yang sama. Pasangan yang dipilih pada validasi juga tidak selalu sama dengan model teratas pada " + T("acc") + ". "
      "Pada Azure 2021 resolusi 5 menit, misalnya, pasangan yang diuji adalah TimesFM-3 dan LSTM, sedangkan model teratas pada "
      "periode uji adalah Toto.")
    ltable("floors", ["Perlakuan galat nol", *DS_LABEL], floor_rows(),
           "Sensitivitas terhadap perlakuan galat nol: nilai TSFM terbaik / pembanding terbaik pada horizon terpanjang. Model terbaik "
           "dipilih ulang untuk setiap baris.", [3.4, 2.1, 2.1, 2.1, 2.1, 2.1, 2.1], size=7)
    P(T("floors") + " memperlihatkan bahwa arah hasil bertahan pada setiap perlakuan galat nol yang masuk akal. Pada empat kombinasi "
      "data padat, nilai TSFM hampir sama untuk batas bawah 10⁻⁹ sampai 0,01, dan TSFM tetap lebih baik pada semua baris. Pada Azure 2021, TSFM tetap lebih baik dengan batas bawah 0,001 sampai 0,1 dan "
      "menurut rasio jumlah galat, tetapi menurut median rasio selisihnya hampir nol (0,500 dan 0,502 pada resolusi 5 menit; 0,506 "
      "dan 0,511 pada 1 menit). Hasil berbalik hanya dengan batas bawah 10⁻⁹. Pada nilai itu, delapan deret dengan MAE model tepat nol "
      "masing-masing menyumbang log rasio sekitar −20,7, sehingga rata-rata geometrik ditentukan oleh segelintir deret. Keunggulan "
      "TSFM pada Azure 2021 dengan demikian nyata tetapi kecil, dan didorong terutama oleh aplikasi yang aktif.")
    ltable("friedman", ["Dataset", "n", "k", "χ²", "p", "CD", "Peringkat 1", "Peringkat 2"], friedman_rows(),
           "Uji Friedman (dengan koreksi nilai seri) dan Nemenyi pada horizon terpanjang. n adalah jumlah deret yang diuji setelah "
           "deret dengan MAE _seasonal naive_ nol dikeluarkan, dan k adalah jumlah model. Nilai p dihitung dengan presisi tinggi dan "
           "dibulatkan ke pangkat sepuluh terdekat.", [2.2, 1.0, 0.9, 1.6, 1.4, 1.1, 3.4, 3.4], size=7.5)
    lfigure("cd", FIG / "cd_azure2019_1min_id.png",
            "Diagram _critical difference_ (CD) dengan uji Nemenyi (α = 0,05) untuk Azure 2019 resolusi 1 menit pada horizon 60 "
            "menit. Angka menunjukkan peringkat rata-rata MAE per deret (1 = terbaik), dan garis horizontal menghubungkan model yang "
            "tidak berbeda secara signifikan. Biru menandai TSFM, oranye menandai model pembanding terlatih, dan hijau menandai model "
            "pembanding statistik.")
    P(T("friedman") + " merangkum uji signifikansi. Uji Friedman menolak hipotesis bahwa semua model setara pada keenam kombinasi, "
      "dengan nilai p antara 10⁻¹¹⁹ dan 10⁻⁹⁵⁵. Meski demikian, model peringkat pertama tidak pernah berbeda secara signifikan "
      "dari model peringkat kedua. Pada Azure 2019 resolusi 1 menit (" + G("cd") + "), TimesFM-3 (3,55) dan Toto (3,74) berada "
      "dalam satu kelompok, sedangkan pembanding terbaik, yaitu N-BEATS (8,17), terpisah lebih dari satu nilai CD (1,21). Karena "
      "itu kami memakai istilah “kelompok teratas” untuk ketiga TSFM terbaik. Pada Azure 2021 resolusi 1 menit, model yang "
      "selalu meramalkan nol justru menempati peringkat rata-rata terbaik, karena pada banyak deret aktualnya nol dan model ini tidak "
      "pernah meleset sama sekali. Rata-rata geometrik pada " + T("acc") + " tetap menempatkan Toto di depan karena selisih galatnya "
      "besar pada aplikasi yang aktif.")
    ltable("perclass", ["Dataset, horizon", "Kelas (n)", "TSFM terbaik", "Pembanding terbaik"], [
        ["Azure19 5m, 60 mnt", "smooth (209)", "0,454 TimesFM-3", "0,541 N-BEATS"],
        ["", "erratic (140)", "0,455 TimesFM-3", "0,566 N-BEATS"],
        ["", "lumpy (79)", "0,474 TimesFM-3", "0,614 N-BEATS"],
        ["", "intermittent (40)", "0,607 TimesFM-3", "0,927 N-BEATS"],
        ["Azure19 1m, 60 mnt", "smooth (149)", "0,474 TimesFM-3", "0,587 N-BEATS"],
        ["", "erratic (131)", "0,487 Chronos-2", "0,613 N-BEATS"],
        ["", "lumpy (115)", "0,530 Toto", "0,673 LightGBM"],
        ["", "intermittent (82)", "0,896 Toto", "1,099 SWA-3"],
        ["Azure21 1m, 60 mnt", "intermittent (42)", "0,478 Toto", "0,642 LightGBM"],
        ["", "lumpy (35)", "0,212 Toto", "0,254 LSTM"],
    ], "relMASE untuk setiap kelas pola. Deret dengan MAE _seasonal naive_ nol dikeluarkan, sehingga jumlah deret (n) lebih kecil "
       "daripada " + T("dataset") + ".", [3.6, 3.0, 3.4, 3.4])
    P("Analisis per kelas pada " + T("perclass") + " memperlihatkan pola yang sama. TSFM terbaik unggul pada setiap kelas di kedua "
      "trace Azure, termasuk kelas _intermittent_ dan _lumpy_. Kedua trace itu berbeda jauh dalam tingkat kejarangannya. Median "
      "aplikasi Azure 2019 aktif selama 1.172 menit per hari, dan 16% aplikasinya tidak memiliki invokasi pada lebih dari 95% menit. "
      "Median aplikasi Azure 2021 hanya aktif 2 menit per hari, dan 79% aplikasinya tidak memiliki invokasi pada lebih dari 95% "
      "menit.")
    ltable("sparse", ["Model", "relMASE aktif", "Rasio jumlah galat aktif", "Ramalan berlebih (tidak aktif)"], sparse_rows(),
           "Azure 2021 resolusi 1 menit, dipisah menurut aktivitas pada periode uji. Kolom aktif mencakup 69 aplikasi yang dipanggil "
           "setidaknya sekali dan memiliki MAE _seasonal naive_ positif. Kolom terakhir adalah rata-rata ramalan (invokasi per menit) "
           "pada 50 aplikasi yang sama sekali tidak dipanggil selama periode uji, sehingga sama dengan MAE-nya.", [3.6, 2.8, 3.4, 4.2])
    P("Pemisahan pada " + T("sparse") + " menjelaskan perbedaan dua sudut pandang tersebut. Dari 119 aplikasi Azure 2021, 50 (42%) "
      "sama sekali tidak dipanggil selama periode uji, dan 42 (35%) dikeluarkan dari relMASE karena MAE _seasonal naive_-nya nol. "
      "Pada aplikasi yang aktif, TSFM unggul menurut kedua metrik: relMASE Toto 0,495, dibanding 0,613 untuk N-BEATS dan 0,680 untuk "
      "LSTM. Pada aplikasi yang tidak aktif, ramalan berlebih seluruh model kurang dari 0,001 invokasi per menit, sehingga TSFM tidak "
      "terus menghasilkan nilai positif yang berarti.")


def results_replication_context() -> None:
    H("4.2 Evaluasi ulang konfigurasi Toner dkk.", 2)
    P("Untuk melihat pengaruh generasi model, kami mengulang protokol Toner dkk. " + c("toner") + " pada data Huawei dengan konteks "
      "520 langkah, horizon 30, 96, dan 336 langkah, serta metrik MASE rata-rata (" + T("toner") + "). Data, panjang konteks, dan "
      "horizonnya sama, tetapi pembagian data uji, titik awal, dan cara agregasinya mengikuti protokol kami, sehingga angka "
      "mutlaknya tidak dapat dibandingkan langsung dengan tabel pada paper tersebut.")
    ltable("toner", ["Model", "H = 30", "H = 96", "H = 336"], [
        ["TimesFM-3", "0,912", "1,102", "1,347"], ["Chronos-2", "0,920", "1,082", "1,301"],
        ["TimesFM-2.5", "0,996", "1,216", "1,453"], ["Toto", "1,004", "1,347", "1,671"],
        ["TiRex", "1,100", "1,402", "1,606"], ["Moirai-2-S", "1,124", "1,558", "2,096"],
        ["Chronos-Bolt-S", "1,222", "1,468", "1,653"], ["TTM-r2", "1,290", "1,600", "–"],
        ["Seasonal naive", "1,324", "1,358", "1,459"], ["Naive", "1,936", "3,413", "4,054"],
    ], "Evaluasi ulang dengan konfigurasi Toner dkk. (Huawei 5 menit, C = 520, MASE rata-rata). Nilai lebih kecil lebih baik.",
       [3.6, 2.4, 2.4, 2.4], bold_cells={(0, 1), (1, 2), (1, 3)})
    P("Chronos-2 dan TimesFM-3 mengalahkan _seasonal naive_ pada ketiga horizon. Model kecil, yaitu Chronos-Bolt-S, Moirai-2-S, dan "
      "TTM-r2, kalah atau setara pada horizon 96 langkah ke atas, sama dengan temuan Toner dkk. Hasil ini mengindikasikan bahwa "
      "hasil negatif pada studi tersebut berkaitan dengan generasi dan ukuran model yang diuji. Atribusi sebab yang lebih kuat "
      "memerlukan model lama dan baru yang dijalankan pada protokol yang identik.")

    H("4.3 Panjang konteks", 2)
    lfigure("context", FIG / "context_ablation_id.png",
            "relMASE pada horizon terpanjang sebagai fungsi panjang konteks C. Panel kiri menunjukkan Azure 2019 resolusi 1 menit, "
            "dan panel kanan menunjukkan Huawei resolusi 5 menit.", 14.5)
    P(G("context") + " memperlihatkan bahwa konteks yang lebih panjang membantu sampai sekitar 2048 langkah pada kedua dataset, dan "
      "sesudah itu perbaikannya kecil. Pada Huawei 5 menit, memperpanjang konteks dari 512 ke 1024 langkah (dari 1,8 ke 3,5 hari) "
      "menurunkan relMASE Chronos-2 dari 0,561 menjadi 0,470, dan perubahan dari 1024 ke 2048 langkah hanya 0,003. Di atas 2048, "
      "TimesFM-3 masih membaik dari 0,471 menjadi 0,436 pada 6912 langkah, Chronos-2 hanya turun ke 0,458, dan Toto sedikit "
      "memburuk (0,515 menjadi 0,523 pada 4096). Pada Azure 2019 1 menit, konteks 2048 langkah (sekitar 34 jam) lebih baik "
      "daripada 512 langkah untuk ketiga model: Chronos-2 dari 0,612 menjadi 0,584, TimesFM-3 dari 0,625 menjadi 0,566, dan Toto "
      "dari 0,588 menjadi 0,556. Di atas 2048, Toto masih sedikit membaik (0,548 pada 4096), sedangkan Chronos-2 sedikit memburuk "
      "(0,596 pada 8192) dan menjadi 4,4 kali lebih lambat. Hasil TiRex identik untuk C ≥ 2048, yang mengindikasikan bahwa "
      "konteksnya dipotong secara internal. Konteks utama C = 2048 dengan demikian wajar untuk kedua resolusi, walaupun TimesFM-3 "
      "pada data 5 menit masih bisa diperbaiki sekitar 7% dengan konteks yang lebih panjang.")


def results_cost_decision() -> None:
    H("4.4 Biaya komputasi (RQ2)", 2)
    ltable("latency", ["Model", "b=1, C=512", "b=1, C=2048", "b=64, C=512", "b=64, C=2048", "RSS (MB)"], [
        ["Naive, seasonal naive", "<0,1", "<0,1", "<0,1", "<0,1", "1300"],
        ["TSB, Croston", "0,2", "0,3", "8,5", "14", "1400"],
        ["TTM-r2", "1,3", "1,7", "6,7", "15", "1586"],
        ["Chronos-Bolt-S", "8,2", "13,5", "94", "338", "1943"],
        ["Moirai-2-S", "4,2", "10,5", "112", "473", "1667"],
        ["Toto (rata-rata)", "27,7", "29,2", "184", "657", "2925"],
        ["Chronos-2", "27,5", "49,5", "477", "1696", "2481"],
        ["ETS (per deret)", "2,3", "8,4", "595", "3409", "1342"],
        ["TimesFM-3", "74,8", "128,7", "1615", "3338", "5044"],
        ["TimesFM-2.5ᵃ", "836", "3348", "838", "3604", "4912"],
        ["TiRex", "223", "223", "9909", "9898", "1887"],
        ["Toto (20 sampel)", "111", "283", "4154", "22994", "9567"],
    ], "Latensi inferensi p50 (ms) di CPU dengan 4 utas dan horizon 60, beserta memori puncak. ᵃ Model dikompilasi dengan _batch_ "
       "internal 64, sehingga angka untuk _batch_ 1 ikut terpengaruh _padding_. RSS (_resident set size_) mencakup sekitar 1,3 GB "
       "data yang dimuat oleh program pengukur.", [3.6, 1.9, 1.9, 1.9, 2.0, 1.7], size=8)
    lfigure("pareto", FIG / "pareto_accuracy_latency_id.png",
            "Akurasi (rata-rata geometrik relMASE pada empat dataset padat) terhadap latensi CPU (_batch_ 64, C = 2048). Garis "
            "titik-titik menunjukkan akurasi model pembanding terlatih yang latensinya tidak diukur.", 10.5)
    P("Pada empat dataset padat, rata-rata geometrik relMASE TimesFM-3 adalah 0,524, Chronos-2 0,534, dan Toto 0,545 (" +
      G("pareto") + "). TimesFM-3 memiliki galat agregat terendah, tetapi membutuhkan 3,3 detik untuk 64 deret, sedangkan Chronos-2 "
      "membutuhkan 1,7 detik dan Toto dalam mode rata-rata 0,66 detik. Kedua model terakhir menawarkan kompromi akurasi dan latensi "
      "yang menarik, dan ketiganya berada pada batas Pareto. TiRex paling lambat di antara TSFM besar (9,9 detik), dan akurasinya "
      "juga lebih rendah. Model ETS yang ditaksir untuk setiap deret pada konteks 2048 sudah sama lambatnya dengan TimesFM-3, "
      "sehingga inferensi TSFM besar tidak lebih mahal daripada penaksiran model statistik lokal. Pada eksekusi penuh, satu siklus "
      "peramalan per jam untuk 550 aplikasi membutuhkan sekitar 37 detik dengan TimesFM-3 dan sekitar 100 detik dengan TiRex di CPU "
      "laptop, jauh di bawah anggaran 60 menit. Sebagai pembanding, pelatihan model untuk Azure 2019 resolusi 1 menit membutuhkan "
      "384 detik (LightGBM), 1615 detik (LSTM), dan 3057 detik (N-BEATS), dan pelatihan itu perlu diulang ketika pola beban berubah.")
    P("Toto dengan 20 sampel 35 kali lebih lambat dan memakai RAM 9,6 GB. Ramalan titik yang diturunkan dari 20 sampel memiliki "
      "relMASE lebih buruk daripada mode rata-rata (Huawei 0,595 dan 0,515; Alibaba 0,697 dan 0,639), tetapi kualitas "
      "probabilistiknya lebih baik atau setara: _weighted quantile loss_ pada Huawei 0,108, dibanding 0,121 untuk mode rata-rata "
      "dengan kuantil konformal, dan pada Alibaba 0,106 dibanding 0,105.")

    H("4.5 Keputusan penskalaan (RQ3)", 2)
    lfigure("tradeoff", FIG / "tradeoff_autoscaling_id.png",
            "Kurva kompromi penskalaan dalam skala logaritmik. Setiap titik pada kurva prediktif dan hibrida adalah pasangan (τ, θ) "
            "yang Pareto-optimal pada periode validasi, dievaluasi pada periode uji. Kurva reaktif memvariasikan durasi _keep-alive_ "
            "atau target utilisasi. Posisi kiri bawah lebih baik.")
    ltable("serverless", ["Kebijakan", "Proksi cold start 2%", "5%", "4,5% (setara adaptasi histogram)"], [
        ["Reaktif (keep-alive)", "0,800", "0,150", "0,186"],
        ["Adaptasi histogram hibrida", "–", "–", "0,199"],
        ["Prediktif Chronos-2", "0,140", "0,047", "0,055"],
        ["Prediktif TimesFM-3", "0,146", "0,048", "0,054"],
        ["Prediktif Toto", "–", "0,062", "0,069"],
        ["Prediktif LightGBM", "–", "0,090", "0,144"],
        ["Prediktif N-BEATS", "–", "0,242", "0,279"],
        ["Prediktif LSTM", "–", "0,331", "–"],
        ["Prediktif naive", "–", "0,674", "0,750"],
    ], "Simulasi serverless Azure 2019 (548 aplikasi, resolusi 1 menit, 2 hari data uji): proksi instance menganggur per kebutuhan "
       "pada proksi _cold start_ yang sama (interpolasi). Tanda “–” berarti tingkat risiko tersebut berada di luar jangkauan "
       "kurva kebijakan.", [4.6, 2.6, 2.0, 4.2], bold_cells={(2, 1), (2, 2), (3, 3)})
    P("**Serverless Azure 2019.** Rencana pemanasan awal berbasis TSFM lebih baik daripada _keep-alive_ reaktif dan model pembanding "
      "di sepanjang kurva (" + T("serverless") + ", panel kiri " + G("tradeoff") + "). Pada proksi _cold start_ 5%, Chronos-2 "
      "membutuhkan 0,047 instance menganggur per kebutuhan, sedangkan kebijakan reaktif membutuhkan 0,150 dan LightGBM 0,090, atau "
      "penghematan sekitar 69% dan 48%. N-BEATS, pembanding yang paling akurat pada Azure 2019, justru membutuhkan 0,242 pada tingkat "
      "yang sama. Pada proksi _cold start_ 2%, kebijakan reaktif membutuhkan 0,800, sedangkan Chronos-2 hanya 0,140, dan kurva "
      "prediktif model pembanding tidak mencapai tingkat risiko serendah itu. Adaptasi histogram hibrida menghasilkan proksi _cold start_ 4,5% dengan "
      "0,199 instance menganggur, sedikit di atas kurva _keep-alive_ reaktif (0,186). Pada tingkat yang sama, Chronos-2 membutuhkan "
      "0,055. Kebijakan hibrida Chronos-2 (τ = 0,9, θ = 0) mencapai proksi _cold start_ 0,50% dengan 0,49 instance menganggur, "
      "sedangkan kebijakan reaktif membutuhkan sekitar 1,05 pada tingkat itu.")
    P("**Serverless Azure 2021.** Pada trace ini manfaatnya tidak konsisten. Rencana prediktif tanpa komponen reaktif tidak pernah "
      "mencapai proksi _cold start_ di bawah 11%, untuk model mana pun. Dalam kebijakan hibrida, TimesFM-3 dan Chronos-2 membutuhkan "
      "2,70 dan 2,71 instance menganggur per kebutuhan pada proksi _cold start_ 10%, lebih sedikit daripada kebijakan reaktif (3,44). "
      "Pada tingkat 7,6%, keadaannya berbalik: TimesFM-3 membutuhkan 5,69, sedangkan kebijakan reaktif 5,08. Kebijakan hibrida "
      "LightGBM, LSTM, dan _naive_ hanya mencapai proksi _cold start_ 10,4% sampai 11,9%. Pada tingkat itu kapasitas menganggur "
      "LightGBM dan _naive_ hampir sama dengan kebijakan reaktif (selisih 2 sampai 3%), sedangkan LSTM 11% lebih hemat. Adaptasi histogram hibrida menghasilkan proksi _cold start_ 7,6% dengan "
      "7,99 instance menganggur per kebutuhan, karena sebagian besar aplikasi tidak memiliki histogram yang representatif dan kembali "
      "ke _keep-alive_ 4 jam.")
    ltable("vm", ["Kebijakan", "Risiko 5%", "Risiko 10%", "Risiko 20%"], [
        ["HPA reaktif", "0,242", "0,187", "0,105"], ["Hibrida TimesFM-3", "0,201", "0,144", "–"],
        ["Hibrida Chronos-2", "–", "0,141", "–"], ["Hibrida MSTL-ARIMA", "–", "0,142", "–"],
        ["Hibrida LightGBM", "–", "0,149", "–"], ["Hibrida N-BEATS", "–", "0,149", "–"],
        ["Prediktif Chronos-2 / TimesFM-3", "–", "0,168", "0,106"], ["Prediktif MSTL-ARIMA", "–", "–", "0,114"],
    ], "Simulasi VM Alibaba: replika berlebih per replika yang dibutuhkan pada risiko kekurangan kapasitas yang sama (hasil "
       "interpolasi). Tanda “–” berarti nilai tersebut berada di luar jangkauan parameter pengatur kebijakan.",
       [6.6, 2.0, 2.0, 2.0], bold_cells={(1, 1), (2, 2)})
    P("**VM Alibaba.** Rencana prediktif per jam tidak lebih baik daripada HPA reaktif (0,106 dan 0,105 pada risiko 20%; " +
      T("vm") + "). Beban VM memiliki autokorelasi yang tinggi, sehingga reaksi setiap 5 menit sudah menangkap sebagian besar "
      "informasi yang dibutuhkan. Kebijakan hibrida menghemat sekitar 24% kapasitas berlebih pada risiko 10% (0,141 dibanding "
      "0,187), tetapi TSFM hampir setara dengan MSTL-ARIMA (0,141 dan 0,142).")
    P("**Apakah urutan akurasi menentukan urutan keputusan?** Korelasi Spearman antara peringkat relMASE dan peringkat biaya pada "
      "risiko yang sama tergolong tinggi, yaitu 0,89 (Alibaba hibrida, risiko 10%), 0,99 (Alibaba prediktif, risiko 20%), dan 0,87 "
      "(Azure 2019 prediktif, proksi _cold start_ 5%). Model yang lebih akurat cenderung menghasilkan keputusan yang lebih baik, "
      "sehingga hipotesis awal kami bahwa urutan model akan banyak berubah tidak terbukti. Besar selisihnya yang berubah. Untuk "
      "membandingkan akurasi dan keputusan secara adil, kami memakai pasangan model yang sama di kedua sisi. Pada kebijakan hibrida "
      "Alibaba, galat Chronos-2 6% lebih kecil daripada MSTL-ARIMA (0,614 dan 0,653), dan kapasitas berlebihnya hanya 1% lebih "
      "kecil (0,141 dan 0,142). Pada Azure 2019 resolusi 1 menit, galat Chronos-2 30% lebih kecil daripada LightGBM (0,584 dan "
      "0,832), dan instance menganggurnya 48% lebih sedikit (0,047 dan 0,090). Terhadap N-BEATS, galatnya 20% lebih kecil (0,584 dan "
      "0,734), sedangkan instance menganggurnya 81% lebih sedikit (0,047 dan 0,242). Pengecualian yang mencolok adalah N-BEATS, yang "
      "lebih akurat daripada LightGBM tetapi menghasilkan keputusan yang lebih buruk.")


def discussion() -> None:
    H("5. Diskusi", 1)
    H("5.1 Kapan TSFM layak dipakai untuk penskalaan otomatis", 2)
    P("Dari hasil di atas kami merumuskan tiga panduan praktis.")
    bullets([
        "**Pemanasan awal untuk aplikasi serverless yang aktif secara teratur.** Pada kasus inilah TSFM paling bermanfaat. Tanpa "
        "pelatihan per aplikasi, rencana pemanasan awal per jam mengurangi proksi instance menganggur hampir separuhnya dibanding "
        "LightGBM, dan lebih dari separuhnya dibanding _keep-alive_.",
        "**Aplikasi yang hampir selalu menganggur.** Ramalan TSFM pada aplikasi seperti ini lebih akurat daripada model pembanding, "
        "tetapi manfaat keputusannya "
        "bergantung pada tingkat risiko yang dipilih operator. Aturan sederhana berdasarkan proporsi menit tanpa invokasi pada "
        "riwayat aplikasi dapat dipakai untuk memutuskan apakah suatu aplikasi cukup ditangani dengan _keep-alive_.",
        "**Penskalaan VM yang sudah reaktif.** Ramalan lebih berguna sebagai batas bawah kapasitas dalam kebijakan hibrida daripada "
        "sebagai pengganti kontrol reaktif. Pada mode ini, model statistik klasik sudah memberikan sebagian besar manfaatnya.",
    ])
    P("Dari sisi biaya, TimesFM-3 dipilih bila galat terendah lebih penting, sedangkan Chronos-2 dan Toto dalam mode rata-rata "
      "dipilih bila latensi di CPU lebih penting. Konteks 2048 langkah memberi sebagian besar manfaat konteks panjang pada kedua "
      "resolusi. Konteks yang lebih panjang lagi hanya layak dicoba untuk TimesFM-3 pada data 5 menit, dengan biaya latensi yang "
      "lebih besar.")
    H("5.2 Kaitan dengan penskalaan hemat energi", 2)
    P("Simulator kami menghasilkan deret kapasitas aktif per menit untuk setiap kebijakan. Deret tersebut dapat langsung dipetakan "
      "ke konsumsi energi dengan model daya linear P = P₀ + (P₁ − P₀)·u, dengan P₀ daya saat menganggur dan P₁ daya pada utilisasi "
      "penuh. Dengan cara ini, hasil RQ3 dapat menjadi masukan bagi studi penskalaan hemat energi yang memakai komponen peramalan "
      "yang sama. Biaya inferensi dari RQ2 juga perlu diperhitungkan sebagai beban energi tambahan dari model peramal itu sendiri.")
    H("5.3 Ancaman terhadap validitas", 2)
    bullets([
        "**Kebocoran data pelatihan awal.** Korpus pelatihan awal lima TSFM kemungkinan besar memuat trace Alibaba 2018 pada tingkat "
        "klaster dan periode yang sama (" + T("audit") + "). Tanpa pembanding yang bebas kontaminasi, besar pengaruhnya terhadap hasil "
        "Alibaba tidak dapat kami pisahkan, sehingga kesimpulan utama kami bersandar pada Azure dan Huawei. Ketepatan audit ini "
        "bergantung pada kelengkapan dokumentasi setiap model, dan korpus TTM-r2 tidak dapat diperiksa.",
        "**Proksi cold start.** Simulator membandingkan kebutuhan dan kapasitas per menit tanpa melacak instance yang sibuk, tanpa "
        "model kedatangan di dalam satu menit, dan tanpa waktu _cold start_ itu sendiri. Angkanya kami tafsirkan sebagai estimasi "
        "relatif antarkebijakan. Jumlah permintaan yang benar-benar mengalami _cold start_ memerlukan simulator tingkat permintaan.",
        "**Rencana per jam.** Kebijakan prediktif memakai ramalan yang dibuat sekali setiap jam. Ramalan yang diperbarui pada setiap "
        "langkah kontrol mungkin mengubah perbandingannya dengan kebijakan reaktif, terutama untuk VM.",
        "**Pemilihan parameter pada validasi.** Pasangan (τ, θ) dipilih pada periode validasi yang juga dipakai untuk mengkalibrasi "
        "kuantil konformal model titik, sehingga pemilihan bagi model tersebut sedikit optimistis.",
        "**Satu mesin.** Seluruh latensi diukur pada satu laptop berprosesor ARM. Angka absolutnya akan berbeda pada server x86, "
        "meskipun urutan relatifnya kemungkinan tidak banyak berubah.",
        "**Konfigurasi model pembanding.** Model terlatih memakai hiperparameter bawaan dengan _early stopping_. LSTM akhirnya "
        "memakai jendela masukan 2 jam, karena konfigurasi awal dengan masukan 864 langkah membutuhkan 3,5 jam pelatihan di CPU dan "
        "tetap menjadi model terburuk. N-BEATS memakai normalisasi standar, karena normalisasi _robust_ gagal pada deret yang jarang "
        "(relMASE validasi Azure 2019 sebesar 1,40, dibanding 0,69 dengan normalisasi standar). Pilihan ini ditentukan dari data "
        "validasi saja, lalu diterapkan ke semua dataset.",
        "**Cakupan data dan statistik.** Penelitian ini hanya memakai empat trace dengan periode uji 2 hingga 4 hari, dan trace Google "
        "2019 belum dievaluasi. Interval _bootstrap_ atas deret tidak mencakup variasi antarperiode maupun korelasi antarderet. "
        "_Block bootstrap_ atas waktu belum kami lakukan.",
        "**Adaptasi histogram hibrida.** Implementasi kami tidak memuat komponen ARIMA dan pembaruan daring dari versi asli Shahrad "
        "dkk., sehingga hasilnya tidak mewakili kinerja metode aslinya secara penuh.",
    ])


def conclusion() -> None:
    H("6. Kesimpulan", 1)
    P("Kami mengevaluasi delapan TSFM secara _zero-shot_ pada empat trace cloud dan serverless publik dari sisi akurasi, biaya CPU, "
      "dan keputusan penskalaan. Galat TSFM generasi 2025 hingga 2026 lebih kecil 7 hingga 26% daripada model pembanding terbaik pada "
      "keenam kombinasi dataset dan resolusi, seluruhnya signifikan menurut uji berpasangan, dan hasil negatif studi sebelumnya tidak "
      "terulang pada model yang lebih baru. Pada aplikasi yang sangat jarang dipanggil, keunggulan itu kecil dan peka terhadap cara "
      "galat nol diperlakukan. Pada tingkat keputusan, manfaatnya besar untuk pemanasan awal instance pada Azure 2019, "
      "tidak konsisten pada aplikasi Azure 2021 yang sangat jarang dipanggil, dan kecil untuk VM yang sudah memakai kontrol reaktif. "
      "Penelitian lanjutan akan mencakup simulator tingkat permintaan, ramalan yang diperbarui pada setiap langkah kontrol, "
      "pemilihan kebijakan secara otomatis untuk setiap aplikasi, dan integrasi dengan model energi.")
    H("Ketersediaan kode dan data", 1)
    P("Paket keterulangan tersedia di https://github.com/harrisonhys/Workload-forecasting dengan lisensi MIT, dan versi yang "
      "dipakai dalam naskah ini (v0.5) diarsipkan di Zenodo dengan DOI https://doi.org/10.5281/zenodo.23662411. Paket ini memuat seluruh kode (persiapan data, pembungkus model dengan revisi _checkpoint_ yang terkunci, "
      "model pembanding, simulator, dan analisis), lingkungan Python dengan berkas kunci versi, serta skrip yang menjalankan "
      "seluruh eksperimen secara berurutan. Pembagian data dicatat pada configs/splits.json, yang memuat batas latih, validasi, "
      "dan uji, jumlah origin, horizon, serta sidik sha256 setiap panel olahan, sehingga pembaca dapat memastikan bahwa data "
      "yang mereka bangun ulang identik. Daftar ID deret untuk setiap dataset dan resolusi ada di configs/series, dan revisi "
      "_checkpoint_ serta pengaturan model pembanding ada di configs/models.json. Keluaran pendukung mencakup metrik per deret "
      "untuk setiap model, hasil simulasi, pengukuran latensi, dan berkas CSV untuk setiap tabel dan gambar. Ramalan per origin "
      "(sekitar 1 GB) dilampirkan pada rilis v0.5 di GitHub dalam enam berkas zip dengan checksum SHA-256, dan juga dapat "
      "dibangkitkan ulang dengan skrip yang sama. Semua trace yang dipakai tersedia untuk publik dengan lisensi masing-masing, dan "
      "model Moirai-2.0-R-small berlisensi CC-BY-NC-4.0.")
    H("Lampiran A. Revisi checkpoint", 1)
    ltable("ckpt", ["Model", "Repositori Hugging Face", "Revisi"], [
        ["TimesFM-3", "google/timesfm-3.0-pytorch", "43046b85ec22"],
        ["TimesFM-2.5", "google/timesfm-2.5-200m-pytorch", "1d952420fba8"],
        ["Chronos-2", "amazon/chronos-2", "29ec3766d36d"],
        ["Chronos-Bolt-S", "amazon/chronos-bolt-small", "772f3d25d38a"],
        ["Toto", "Datadog/Toto-Open-Base-1.0", "0411ceb27bdf"],
        ["Moirai-2-S", "Salesforce/moirai-2.0-R-small", "30f43ff08c84"],
        ["TiRex", "NX-AI/TiRex", "63c740922493"],
        ["TTM-r2", "ibm-granite/granite-timeseries-ttm-r2", "branch 1536-96-r2 (f6ed09d67cf4)"],
    ], "Revisi _checkpoint_ yang dipakai (12 karakter pertama hash commit). Varian TTM-r2 dipilih otomatis oleh pustaka granite-tsfm "
       "berdasarkan panjang konteks, dan hash branch dicatat per 26 Oktober 2026.", [2.8, 6.4, 4.8], size=8)


def references() -> None:
    H("Daftar Pustaka", 1)
    for i, k in enumerate(ORDER, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.8), Cm(-0.8)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"[{i}]\t{REFS[k]}")
        r.font.size = Pt(9)
    unused = set(REFS) - set(ORDER)
    if unused:
        raise ValueError(f"references defined but never cited: {sorted(unused)}")


def main() -> None:
    setup()
    for part in (front_matter, intro, related, method, results_accuracy, results_replication_context,
                 results_cost_decision, discussion, conclusion, references):
        part()
    doc.save(OUT)
    print(f"wrote {OUT}  ({N['tab']} tables, {N['fig']} figures, {len(ORDER)} references)")


if __name__ == "__main__":
    main()
